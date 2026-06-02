import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

def create_time_sheet_template(output_path="timesheet_placeholder.docx"):
    # 1. Initialize Document
    doc = Document()

    # 2. Page Setup (Margins)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 3. Configure Base Styles (Fonts)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # 4. Add Document Title (Tätigkeitsnachweis)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(18)
    title_run = title_p.add_run("Tätigkeitsnachweis")
    title_run.font.size = Pt(22)
    title_run.font.bold = True

    # 5. Add Employee Metadata Fields
    metadata = [
        ("Name Mitarbeiter /-in:", "Max Mustermann"),
        ("Personal Nummer:", "12345678"),
        ("Stadt:", "Istanbul")
    ]
    
    for label, placeholder in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        label_run = p.add_run(f"{label} ")
        label_run.bold = True
        
        value_run = p.add_run(placeholder)
        value_run.underline = True

    # Add spacing before table
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)

    # 6. Create Time Sheet Table
    # 1 header row + 31 data rows = 32 total rows
    headers = ["Datum", "Arbeitsbeginn", "Arbeitsende", "Pause", "Arbeitszeit", "Sonstiges"]
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'  # Standard visible grid structure

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        # Bold and center header text
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Explicit column widths (Total width ~ 6.9 inches)
    col_widths = [Inches(1.2), Inches(1.1), Inches(1.1), Inches(0.9), Inches(1.1), Inches(1.5)]

    # Mock Data Generation for 31 days (January 2026 as a placeholder)
    # Alternates between simulated workdays and empty weekend days to show layout flexibility
    for day in range(1, 32):
        row_cells = table.add_row().cells
        date_str = f"{day:02d}.01.2026"
        
        # Simple rule to simulate weekends off for placeholder visualization
        is_weekend = day in [3, 4, 10, 11, 17, 18, 24, 25, 31]
        
        if is_weekend:
            row_data = [date_str, "", "", "", "", "Wochenende"]
        else:
            row_data = [date_str, "08:15", "14:45", "00:30", "6,0", "Präsenz"]

        # Populate and style cells
        for i, text in enumerate(row_data):
            row_cells[i].text = text
            p = row_cells[i].paragraphs[0]
            
            # Align center for times/dates, left-align for notes
            if i < 5:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Apply column widths across all rows systematically
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # 7. Total Summary Block
    summary_p = doc.add_paragraph()
    summary_p.paragraph_format.space_before = Pt(16)
    summary_p.paragraph_format.space_after = Pt(40)
    
    summary_run_label = summary_p.add_run("Summe der Arbeitsstunden im Monat: ")
    summary_run_label.bold = True
    
    summary_run_val = summary_p.add_run("132,0 Stunden")
    summary_run_val.underline = True

    # 8. Footer Sign-off Section (Datum, Unterschrift)
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(24)
    
    # Left side: Date line
    date_line = footer_p.add_run("Datum: ______________________")
    
    # Add tab spaces to separate cleanly on the same line
    footer_p.add_run("\t\t\t\t")
    
    # Right side: Signature line
    sig_line = footer_p.add_run("Unterschrift: ______________________")
    
    # Small caption underneath lines
    caption_p = doc.add_paragraph()
    caption_p.paragraph_format.space_before = Pt(2)
    caption_run = caption_p.add_run("Datum, Unterschrift")
    caption_run.font.size = Pt(9)
    caption_run.font.italic = True

    # Save Document
    doc.save(output_path)
    print(f"Success! Process completed. File saved to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_time_sheet_template()