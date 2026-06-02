import os
import tempfile
import pytest
from docx import Document
import exporter.docx_exporter as docx_exporter

def test_generate_docx_creation():
    # Arrange data
    employee_name = "Jane Doe"
    personal_id = "P112233"
    city_name = "Istanbul"
    year = 2026
    month = 6
    target_hours = 80.0
    
    daily_entries = [
        {"work_date": "2026-06-01", "hours_worked": 8.0, "start_time": "08:00", "end_time": "16:30", "break_duration": "00:30", "remarks": "Shift 1"},
        {"work_date": "2026-06-02", "hours_worked": 0.0, "start_time": "", "end_time": "", "break_duration": "", "remarks": ""}, # Off-day (weekday)
        {"work_date": "2026-06-03", "hours_worked": 4.5, "start_time": "12:00", "end_time": "16:30", "break_duration": "00:00", "remarks": "Shift 2"},
        {"work_date": "2026-06-07", "hours_worked": 0.0, "start_time": "", "end_time": "", "break_duration": "", "remarks": ""}  # Weekend off-day
    ]
    
    # Create a temporary file path
    temp_fd, temp_path = tempfile.mkstemp(suffix=".docx")
    os.close(temp_fd)
    
    try:
        # Act
        exported_path = docx_exporter.generate_docx(
            employee_name=employee_name,
            personal_id=personal_id,
            city_name=city_name,
            year=year,
            month=month,
            daily_entries=daily_entries,
            target_hours=target_hours,
            output_path=temp_path
        )
        
        # Assert
        assert os.path.exists(exported_path)
        
        # 1. Parse using python-docx
        doc = Document(exported_path)
        
        # 2. Check title
        assert doc.paragraphs[0].text == "Tätigkeitsnachweis"
        
        # 3. Check metadata
        metadata_text = [p.text for p in doc.paragraphs[1:4]]
        assert "Name Mitarbeiter /-in: Jane Doe" in metadata_text
        assert "Personal Nummer: P112233" in metadata_text
        assert "Stadt: Istanbul" in metadata_text
        
        # 4. Check table structure
        assert len(doc.tables) == 1
        table = doc.tables[0]
        
        # Header + 4 daily entries = 5 rows
        assert len(table.rows) == 5
        
        # Check header labels
        headers = [cell.text for cell in table.rows[0].cells]
        assert headers == ["Datum", "Arbeitsbeginn", "Arbeitsende", "Pause", "Arbeitszeit", "Sonstiges"]
        
        # Row 1 check (June 1)
        r1_cells = [cell.text for cell in table.rows[1].cells]
        assert r1_cells[0] == "01.06.2026"
        assert r1_cells[1] == "08:00"
        assert r1_cells[2] == "16:30"
        assert r1_cells[3] == "00:30"
        assert r1_cells[4] == "8"
        assert r1_cells[5] == "Shift 1"
        
        # Row 2 check (June 2 - Off-day)
        r2_cells = [cell.text for cell in table.rows[2].cells]
        assert r2_cells[0] == "02.06.2026"
        assert r2_cells[1] == ""
        assert r2_cells[2] == ""
        assert r2_cells[3] == ""
        assert r2_cells[4] == ""
        assert r2_cells[5] == ""
        
        # Row 3 check (June 3 - Half hour)
        r3_cells = [cell.text for cell in table.rows[3].cells]
        assert r3_cells[0] == "03.06.2026"
        assert r3_cells[1] == "12:00"
        assert r3_cells[2] == "16:30"
        assert r3_cells[3] == "" # 00:00 or empty should be empty
        assert r3_cells[4] == "4,5" # German comma
        assert r3_cells[5] == "Shift 2"
        
        # Row 4 check (June 7 - Weekend off-day)
        r4_cells = [cell.text for cell in table.rows[4].cells]
        assert r4_cells[0] == "07.06.2026"
        assert r4_cells[1] == ""
        assert r4_cells[2] == ""
        assert r4_cells[3] == ""
        assert r4_cells[4] == ""
        assert r4_cells[5] != "Wochenende" # Ensure NO weekend label
        assert r4_cells[5] == ""
        
        # 5. Check summary sum
        # total hours is 8.0 + 4.5 = 12.5. formatted as 12,5
        summary_text = doc.paragraphs[-3].text
        assert "Summe der Arbeitsstunden im Monat: 12,5 Stunden" in summary_text
        
    finally:
        # Cleanup
        try:
            os.remove(temp_path)
        except OSError:
            pass
